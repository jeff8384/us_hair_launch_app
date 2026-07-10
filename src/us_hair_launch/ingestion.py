from __future__ import annotations

import csv
from collections.abc import Iterable
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .models import FilePreview, PreviewResult, RawRow, Retailer

SUPPORTED_SUFFIXES = {".csv", ".xlsx"}


def infer_retailer(*parts: str) -> Retailer:
    text = " ".join(parts).lower()
    if "sephora" in text:
        return "sephora"
    if "ulta" in text:
        return "ulta"
    return "unknown"


def scan_input_dir(input_dir: Path) -> PreviewResult:
    diagnostics: list[str] = []
    previews: list[FilePreview] = []
    if not input_dir.exists():
        return PreviewResult(
            input_dir=str(input_dir),
            files=[],
            diagnostics=["Input directory missing"],
        )
    for path in sorted(input_dir.iterdir()):
        if path.name.startswith(".") or path.suffix.lower() not in SUPPORTED_SUFFIXES:
            continue
        try:
            previews.append(preview_file(path))
        except Exception as exc:
            diagnostics.append(f"{path.name}: {type(exc).__name__}: {exc}")
    if not previews:
        diagnostics.append("No supported input files found")
    return PreviewResult(input_dir=str(input_dir), files=previews, diagnostics=diagnostics)


def preview_file(path: Path, sample_size: int = 3) -> FilePreview:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        rows = list(_read_csv_rows(path, limit=sample_size))
        columns = list(rows[0].values) if rows else []
        return FilePreview(
            path=str(path),
            filename=path.name,
            file_type="csv",
            retailer_hint=infer_retailer(path.name),
            sheets=["CSV"],
            columns=columns,
            sample_rows=[row.values for row in rows],
        )
    rows, sheets, columns = _preview_xlsx(path, sample_size)
    return FilePreview(
        path=str(path),
        filename=path.name,
        file_type="xlsx",
        retailer_hint=infer_retailer(path.name, " ".join(sheets)),
        sheets=sheets,
        columns=columns,
        sample_rows=rows,
    )


def read_raw_rows(input_dir: Path) -> tuple[list[RawRow], list[str]]:
    preview = scan_input_dir(input_dir)
    rows: list[RawRow] = []
    diagnostics = list(preview.diagnostics)
    for item in preview.files:
        path = Path(item.path)
        try:
            if path.suffix.lower() == ".csv":
                rows.extend(_read_csv_rows(path))
            else:
                rows.extend(_read_xlsx_rows(path))
        except Exception as exc:
            diagnostics.append(f"{path.name}: {type(exc).__name__}: {exc}")
    return rows, diagnostics


def _read_csv_rows(path: Path, limit: int | None = None) -> Iterable[RawRow]:
    last_error: Exception | None = None
    for encoding in ("utf-8-sig", "utf-8", "cp949", "latin-1"):
        try:
            with path.open("r", encoding=encoding, newline="") as handle:
                reader = csv.DictReader(handle)
                for offset, row in enumerate(reader, start=2):
                    if limit is not None and offset > limit + 1:
                        break
                    yield RawRow(
                        source_file=path.name,
                        source_sheet="CSV",
                        row_index=offset,
                        values={str(k): v for k, v in row.items() if k is not None},
                    )
            return
        except UnicodeDecodeError as exc:
            last_error = exc
    assert last_error is not None
    raise last_error


def _read_xlsx_rows(path: Path) -> Iterable[RawRow]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    for sheet in workbook.worksheets:
        iterator = sheet.iter_rows(values_only=True)
        header = _header(next(iterator, ()))
        for row_index, row in enumerate(iterator, start=2):
            values = {
                header[index]: value for index, value in enumerate(row) if index < len(header)
            }
            if any(value not in (None, "") for value in values.values()):
                yield RawRow(
                    source_file=path.name,
                    source_sheet=sheet.title,
                    row_index=row_index,
                    values=values,
                )


def _preview_xlsx(
    path: Path,
    sample_size: int,
) -> tuple[list[dict[str, Any]], list[str], list[str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets = workbook.sheetnames
    rows: list[dict[str, Any]] = []
    columns: list[str] = []
    for sheet in workbook.worksheets:
        iterator = sheet.iter_rows(values_only=True)
        columns = _header(next(iterator, ()))
        for _, row in zip(range(sample_size), iterator, strict=False):
            rows.append(
                {columns[index]: value for index, value in enumerate(row) if index < len(columns)}
            )
        break
    return rows, sheets, columns


def _header(row: tuple[Any, ...]) -> list[str]:
    used: dict[str, int] = {}
    columns: list[str] = []
    for index, value in enumerate(row):
        base = str(value).strip() if value not in (None, "") else f"column_{index + 1}"
        count = used.get(base, 0)
        used[base] = count + 1
        columns.append(base if count == 0 else f"{base}_{count + 1}")
    return columns
